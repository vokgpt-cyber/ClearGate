"""Test harness for Cleargate anonymization pipeline.

Runs the 2-layer pipeline (Presidio regex + spaCy ru_core_news_sm) on the
7 generated test documents and collects:
- All detected entities with type/score/source_layer
- Anonymized text output (original -> placeholders)
- Deanonymized text (round-trip check)
- Per-doc and aggregate stats

No LLM layer (Ollama not in sandbox).
"""

from __future__ import annotations

import asyncio
import json
import os
import secrets
import sys
import time
from pathlib import Path

# Make backend importable
BACKEND = Path("/sessions/affectionate-pensive-planck/mnt/Velum/backend")
sys.path.insert(0, str(BACKEND))

# spaCy uses ru_core_news_sm in sandbox (ru_core_news_lg download is too large)
os.environ.setdefault("CLEARGATE_SPACY_MODEL", "ru_core_news_sm")

from docx import Document as DocxDocument  # noqa: E402

from app.services.entity_registry import EntityRegistry  # noqa: E402
from app.services.ner_pipeline import NERPipeline  # noqa: E402

TEST_DIR = Path("/sessions/affectionate-pensive-planck/mnt/Velum/test")
RESULTS_DIR = TEST_DIR / "results"
RESULTS_DIR.mkdir(parents=True, exist_ok=True)

DOCS = [
    "01_NDA.docx",
    "02_Dogovor_uslug.docx",
    "03_Trudovoy_dogovor.docx",
    "04_Agentskiy_dogovor.docx",
    "05_Politika_PDn.docx",
    "06_Dop_soglashenie.docx",
    "07_Protokol_sobraniya.txt",
]


def extract_text(path: Path) -> str:
    """Extract plain text from .docx or .txt preserving paragraph order.

    For .docx walk body and include paragraphs + table cell text row by row.
    """
    if path.suffix.lower() == ".txt":
        return path.read_text(encoding="utf-8")

    doc = DocxDocument(str(path))
    chunks: list[str] = []
    # Walk body elements in document order: paragraphs and tables
    for block in doc.element.body.iterchildren():
        tag = block.tag.split("}", 1)[-1]
        if tag == "p":
            # Find the python-docx paragraph wrapper
            for p in doc.paragraphs:
                if p._element is block:
                    if p.text.strip():
                        chunks.append(p.text)
                    break
        elif tag == "tbl":
            for t in doc.tables:
                if t._element is block:
                    for row in t.rows:
                        cells_text = [c.text.strip() for c in row.cells]
                        # Dedup horizontally merged cells (same text repeated)
                        seen = []
                        for ct in cells_text:
                            if not seen or seen[-1] != ct:
                                seen.append(ct)
                        chunks.append(" | ".join(seen))
                    break
    return "\n".join(chunks)


async def analyze_document(pipeline: NERPipeline, path: Path) -> dict:
    """Run pipeline on a single doc, return summary dict."""
    text = extract_text(path)
    n_chars = len(text)
    n_words = len(text.split())

    t0 = time.perf_counter()
    entities = await pipeline.analyze(text, language="ru", score_threshold=0.3)
    elapsed = time.perf_counter() - t0

    # Build registry & anonymize
    registry = EntityRegistry(
        master_key=secrets.token_bytes(32),
        session_id=f"test-{path.stem}",
        locale="ru",
    )
    anonymized = registry.anonymize_text(text, entities)
    deanonymized = registry.deanonymize_text(anonymized)

    # Stats
    by_type: dict[str, int] = {}
    by_layer: dict[str, int] = {}
    for e in entities:
        by_type[e.entity_type] = by_type.get(e.entity_type, 0) + 1
        by_layer[e.source_layer] = by_layer.get(e.source_layer, 0) + 1

    roundtrip_ok = text == deanonymized

    return {
        "doc": path.name,
        "chars": n_chars,
        "words": n_words,
        "entities_total": len(entities),
        "by_type": dict(sorted(by_type.items())),
        "by_layer": dict(sorted(by_layer.items())),
        "elapsed_sec": round(elapsed, 2),
        "roundtrip_ok": roundtrip_ok,
        "registry_entries": registry.entity_count,
        "text": text,
        "anonymized": anonymized,
        "deanonymized": deanonymized,
        "entities": [
            {
                "text": e.text,
                "type": e.entity_type,
                "start": e.start,
                "end": e.end,
                "score": round(e.score, 3),
                "layer": e.source_layer,
            }
            for e in entities
        ],
    }


async def main() -> None:
    print(f"[test-runner] starting, test dir = {TEST_DIR}", flush=True)
    print("[test-runner] building pipeline (spaCy + Presidio, no LLM)...", flush=True)

    pipeline = NERPipeline(
        spacy_model=os.environ["CLEARGATE_SPACY_MODEL"],
        gliner_model=None,   # GLiNER model too large for sandbox
        enable_llm_layer=False,  # Ollama not available
    )
    print("[test-runner] pipeline ready\n", flush=True)

    all_results: list[dict] = []
    for name in DOCS:
        path = TEST_DIR / name
        if not path.exists():
            print(f"  [!] missing {name}", flush=True)
            continue
        print(f"  -> {name} ...", flush=True, end="")
        result = await analyze_document(pipeline, path)
        print(
            f" {result['entities_total']} entities "
            f"({result['elapsed_sec']}s, roundtrip={'OK' if result['roundtrip_ok'] else 'FAIL'})",
            flush=True,
        )
        all_results.append(result)

        # Persist per-doc outputs
        stem = path.stem
        (RESULTS_DIR / f"{stem}_original.txt").write_text(
            result["text"], encoding="utf-8"
        )
        (RESULTS_DIR / f"{stem}_anonymized.txt").write_text(
            result["anonymized"], encoding="utf-8"
        )
        (RESULTS_DIR / f"{stem}_deanonymized.txt").write_text(
            result["deanonymized"], encoding="utf-8"
        )
        (RESULTS_DIR / f"{stem}_entities.json").write_text(
            json.dumps(result["entities"], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    # Aggregate summary
    summary = {
        "total_docs": len(all_results),
        "total_chars": sum(r["chars"] for r in all_results),
        "total_entities": sum(r["entities_total"] for r in all_results),
        "roundtrip_ok": sum(1 for r in all_results if r["roundtrip_ok"]),
        "roundtrip_fail": sum(1 for r in all_results if not r["roundtrip_ok"]),
        "per_doc": [
            {k: v for k, v in r.items() if k not in ("text", "anonymized", "deanonymized", "entities")}
            for r in all_results
        ],
    }
    # Global type counts
    global_by_type: dict[str, int] = {}
    for r in all_results:
        for t, n in r["by_type"].items():
            global_by_type[t] = global_by_type.get(t, 0) + n
    summary["global_by_type"] = dict(sorted(global_by_type.items(), key=lambda kv: -kv[1]))

    (RESULTS_DIR / "_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print("\n[test-runner] DONE.", flush=True)
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    asyncio.run(main())
