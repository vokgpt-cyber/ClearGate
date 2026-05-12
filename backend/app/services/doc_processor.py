"""Document parsing for DOCX, PDF, and TXT files.

Parses uploaded files into plain text in memory (no temp files on disk).
For non-DOCX formats, also builds a conservative DOCX representation so the
existing review, export, and deanonymization flow can operate on one internal
document shape.
"""

from __future__ import annotations

import io
import re
import zipfile
from copy import deepcopy

import structlog
from pydantic import BaseModel, Field

logger = structlog.get_logger(__name__)

_LEGAL_FORM_LINEBREAK = re.compile(
    r"\b(ООО|ОАО|АО|ПАО|ЗАО)\s+([«\"][^»\"\n]{2,80})\n\s*([^»\"]{2,100}[»\"])",
    re.IGNORECASE,
)
_REQUISITE_LABEL_INLINE = re.compile(
    r"\s+((?:Банк|ИНН|ОГРН|КПП|БИК|Адрес|Арендодатель|Арендатор)\s*:)",
    re.IGNORECASE,
)
_PDF_MULTI_SPACE = re.compile(r"[ \t]{2,}")

_MAX_DOCX_ENTRIES = 2_000
_MAX_DOCX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
_MAX_DOCX_COMPRESSION_RATIO = 100


class ParseResult(BaseModel):
    """Result of document parsing."""

    text: str
    page_count: int | None = None
    render_docx_bytes: bytes | None = None
    render_filename: str | None = None
    ocr_used: bool = False
    warnings: list[str] = Field(default_factory=list)


class DocumentProcessor:
    """Parse documents (DOCX, PDF, TXT) to plain text."""

    def parse(self, content: bytes, format: str) -> ParseResult:
        """Parse document bytes into plain text.

        Args:
            content: Raw file bytes.
            format: File format ("docx", "pdf", "txt").

        Returns:
            ParseResult with extracted text and optional page count.

        Raises:
            ValueError: If format is unsupported.
        """
        if format == "txt":
            return self._parse_txt(content)
        if format == "docx":
            return self._parse_docx(content)
        if format == "pdf":
            return self._parse_pdf(content)
        raise ValueError(f"Unsupported format: {format}")

    def combine_results(self, results: list[ParseResult]) -> ParseResult:
        """Combine multiple parsed documents into one internal DOCX shape."""
        if not results:
            raise ValueError("No documents to combine")
        if len(results) == 1:
            return results[0]

        render_parts = [result.render_docx_bytes for result in results]
        if any(part is None for part in render_parts):
            raise ValueError("All documents must have DOCX render bytes")

        text = "\n\n".join(result.text.strip() for result in results if result.text.strip())
        page_counts = [result.page_count for result in results if result.page_count is not None]
        warnings: list[str] = []
        for result in results:
            warnings.extend(result.warnings)

        return ParseResult(
            text=text,
            page_count=sum(page_counts) if page_counts else None,
            render_docx_bytes=self._combine_docx_bytes(
                [part for part in render_parts if part is not None]
            ),
            render_filename="combined_document.docx",
            ocr_used=any(result.ocr_used for result in results),
            warnings=warnings,
        )

    def _parse_txt(self, content: bytes) -> ParseResult:
        """Parse plain text file."""
        text = content.decode("utf-8", errors="replace")
        return ParseResult(
            text=text,
            render_docx_bytes=self._text_to_docx(text),
            render_filename="document.docx",
        )

    def _parse_docx(self, content: bytes) -> ParseResult:
        """Parse DOCX file using python-docx."""
        from docx import Document

        self._validate_docx_zip(content)
        try:
            doc = Document(io.BytesIO(content))
        except Exception as exc:
            raise ValueError("Invalid DOCX file") from exc
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)
        logger.info("doc_processor.docx", paragraphs=len(paragraphs))
        return ParseResult(text=text, render_docx_bytes=content)

    def _validate_docx_zip(self, content: bytes) -> None:
        """Reject malformed or suspiciously compressed DOCX packages."""
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid DOCX file") from exc

        if len(infos) > _MAX_DOCX_ENTRIES:
            raise ValueError("DOCX file has too many internal entries")

        total_uncompressed = 0
        for info in infos:
            filename = info.filename.replace("\\", "/")
            if filename.startswith("/") or ".." in filename.split("/"):
                raise ValueError("DOCX file contains unsafe internal paths")
            if info.is_dir():
                continue
            total_uncompressed += info.file_size
            if total_uncompressed > _MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX file expands to too much data")
            if (
                info.compress_size > 0
                and info.file_size / info.compress_size > _MAX_DOCX_COMPRESSION_RATIO
            ):
                raise ValueError("DOCX file compression ratio is suspicious")

    def _parse_pdf(self, content: bytes) -> ParseResult:
        """Parse PDF file using PyMuPDF, with OCR fallback when available."""
        import pymupdf

        doc = pymupdf.open(stream=content, filetype="pdf")
        pages: list[str] = []
        ocr_used = False
        strikeout_chars_removed = 0
        warnings: list[str] = []
        page_sizes: list[tuple[float, float]] = []

        for idx, page in enumerate(doc, start=1):
            page_sizes.append((float(page.rect.width), float(page.rect.height)))
            text = page.get_text("text", sort=True)
            text, removed = self._extract_pdf_text_without_strikeouts(page, text)
            strikeout_chars_removed += removed
            if not text.strip():
                ocr_text = self._ocr_page_text(page, idx)
                if ocr_text.strip():
                    text = ocr_text
                    ocr_used = True
                else:
                    warnings.append(f"Page {idx} has no extractable text; OCR was unavailable or empty.")
            pages.append(self._normalize_pdf_text(text).rstrip())

        text = "\n".join(pages)
        page_count = len(pages)
        render_docx = self._pdf_pages_to_docx(pages, page_sizes)
        doc.close()
        logger.info(
            "doc_processor.pdf",
            pages=page_count,
            ocr_used=ocr_used,
            strikeout_chars_removed=strikeout_chars_removed,
            warnings=len(warnings),
        )
        return ParseResult(
            text=text,
            page_count=page_count,
            render_docx_bytes=render_docx,
            render_filename="document_from_pdf.docx",
            ocr_used=ocr_used,
            warnings=warnings,
        )

    def _ocr_page_text(self, page, page_number: int) -> str:  # type: ignore[no-untyped-def]
        """Try PyMuPDF OCR for image-only pages.

        This requires Tesseract data in the runtime image. The method is
        intentionally best-effort: a missing OCR engine should not break text
        PDFs or DOCX/TXT uploads.
        """
        try:
            textpage = page.get_textpage_ocr(language="rus+eng", full=True)
            return page.get_text("text", textpage=textpage, sort=True)
        except Exception:
            logger.warning(
                "doc_processor.pdf_ocr_unavailable",
                page=page_number,
                exc_info=True,
            )
            return ""

    def _normalize_pdf_text(self, text: str) -> str:
        """Clean common PDF extraction artifacts before NER.

        Text PDFs often split a quoted company name across visual lines or
        keep requisites such as "Адрес ... Банк ..." on one extracted line.
        The DOCX path does not have those artifacts, so normalize only here.
        """
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        previous = None
        while previous != normalized:
            previous = normalized
            normalized = _LEGAL_FORM_LINEBREAK.sub(
                lambda m: f"{m.group(1)} {m.group(2)} {m.group(3).strip()}",
                normalized,
            )
        normalized = _REQUISITE_LABEL_INLINE.sub(r"\n\1", normalized)
        normalized = re.sub(r"[ \t]+\n", "\n", normalized)
        normalized = re.sub(r"\n{3,}", "\n\n", normalized)
        return normalized

    def _extract_pdf_text_without_strikeouts(self, page, fallback_text: str) -> tuple[str, int]:  # type: ignore[no-untyped-def]
        """Extract text while removing Word Track Changes deletions.

        Word exports deleted text as normal selectable PDF text and draws a
        separate horizontal strike line through it. PyMuPDF therefore exposes
        the deleted words through ``get_text("text")``. We remove only chars
        crossed through the middle of their bbox; underlined inserted text
        stays intact because underline geometry sits near the bbox bottom.
        """
        strike_lines = self._pdf_horizontal_lines(page)
        if not strike_lines:
            return fallback_text, 0

        margin_start = self._pdf_revision_margin_start(page, strike_lines)
        raw = page.get_text("rawdict", sort=True)
        output_lines: list[str] = []
        removed = 0

        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                line_bbox = line.get("bbox")
                if (
                    margin_start is not None
                    and line_bbox is not None
                    and float(line_bbox[0]) >= margin_start
                ):
                    removed += sum(
                        1
                        for span in line.get("spans", [])
                        for char in span.get("chars", [])
                        if char.get("c")
                    )
                    continue
                chars: list[str] = []
                for span in line.get("spans", []):
                    for char in span.get("chars", []):
                        value = char.get("c", "")
                        bbox = char.get("bbox")
                        if not value or bbox is None:
                            continue
                        if self._pdf_char_is_struck(bbox, strike_lines):
                            removed += 1
                            continue
                        chars.append(value)
                line_text = _PDF_MULTI_SPACE.sub(" ", "".join(chars)).rstrip()
                if line_text.strip():
                    output_lines.append(line_text)

        if removed == 0:
            return fallback_text, 0
        return "\n".join(output_lines), removed

    @staticmethod
    def _pdf_horizontal_lines(page) -> list[tuple[float, float, float, float, object | None]]:  # type: ignore[no-untyped-def]
        lines: list[tuple[float, float, float, float, object | None]] = []
        for drawing in page.get_drawings():
            color = drawing.get("color")
            for item in drawing.get("items", []):
                if item[0] == "l":
                    p1, p2 = item[1], item[2]
                    if abs(p1.y - p2.y) > 1.0 or abs(p1.x - p2.x) < 2.0:
                        continue
                    x0, x1 = sorted((float(p1.x), float(p2.x)))
                    y = (float(p1.y) + float(p2.y)) / 2.0
                    lines.append((x0, y, x1, y, color))
                elif item[0] == "re":
                    rect = item[1]
                    if rect.width < 2.0 or rect.height > 2.0:
                        continue
                    y = (float(rect.y0) + float(rect.y1)) / 2.0
                    lines.append((float(rect.x0), y, float(rect.x1), y, color))
        return lines

    @staticmethod
    def _pdf_revision_margin_start(page, lines: list[tuple[float, float, float, float, object | None]]) -> float | None:  # type: ignore[no-untyped-def]
        """Detect and drop Word Track Changes balloons in the right margin."""
        page_width = float(page.rect.width)
        right_third = page_width * 0.65
        candidates = [
            min(x0, x1)
            for x0, _y0, x1, _y1, color in lines
            if color is not None and min(x0, x1) >= right_third and abs(x1 - x0) >= 20
        ]
        if len(candidates) < 2:
            return None
        return min(candidates) - 2.0

    @staticmethod
    def _pdf_char_is_struck(
        bbox: tuple[float, float, float, float],
        lines: list[tuple[float, float, float, float, object | None]],
    ) -> bool:
        x0, y0, x1, y1 = map(float, bbox)
        width = max(x1 - x0, 0.1)
        height = max(y1 - y0, 0.1)
        center_x = (x0 + x1) / 2.0

        for line_x0, line_y, line_x1, _line_y1, _color in lines:
            if center_x < min(line_x0, line_x1) - 0.5:
                continue
            if center_x > max(line_x0, line_x1) + 0.5:
                continue
            vertical_ratio = (line_y - y0) / height
            if not 0.32 <= vertical_ratio <= 0.74:
                continue
            overlap = min(x1, max(line_x0, line_x1)) - max(x0, min(line_x0, line_x1))
            if overlap / width >= 0.25:
                return True
        return False

    def _text_to_docx(self, text: str) -> bytes:
        """Create a simple DOCX wrapper for plain text."""
        return self._pages_to_docx([text], [(595.0, 842.0)])

    def _pdf_pages_to_docx(
        self,
        pages: list[str],
        page_sizes: list[tuple[float, float]],
    ) -> bytes:
        """Convert extracted PDF page text into a reviewable DOCX."""
        return self._pages_to_docx(pages, page_sizes or [(595.0, 842.0)])

    def _combine_docx_bytes(self, documents: list[bytes]) -> bytes:
        """Append DOCX body elements while preserving paragraph/table formatting."""
        from docx import Document

        if not documents:
            raise ValueError("No DOCX documents to combine")
        master = Document(io.BytesIO(documents[0]))
        body = master.element.body

        def insert_before_section(element) -> None:  # type: ignore[no-untyped-def]
            section = body.sectPr
            if section is None:
                body.append(element)
            else:
                body.insert(body.index(section), element)

        for content in documents[1:]:
            master.add_page_break()
            source = Document(io.BytesIO(content))
            for child in source.element.body:
                if child.tag.endswith("}sectPr"):
                    continue
                insert_before_section(deepcopy(child))

        buffer = io.BytesIO()
        master.save(buffer)
        return buffer.getvalue()

    def _pages_to_docx(
        self,
        pages: list[str],
        page_sizes: list[tuple[float, float]],
    ) -> bytes:
        from docx import Document
        from docx.shared import Pt

        document = Document()
        first_width, first_height = page_sizes[0]
        section = document.sections[0]
        section.page_width = Pt(first_width)
        section.page_height = Pt(first_height)
        margin = Pt(36)
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin

        normal = document.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(2)
        normal.paragraph_format.line_spacing = 1.0

        for page_index, page_text in enumerate(pages):
            if page_index > 0:
                document.add_page_break()
            lines = page_text.splitlines() or [""]
            for raw_line in lines:
                paragraph = document.add_paragraph(raw_line.rstrip())
                paragraph.paragraph_format.space_after = Pt(1)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
