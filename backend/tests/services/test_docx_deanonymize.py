"""Tests for DOCX deanonymization — PlaceholderMatcher and deanonymize_docx."""

from __future__ import annotations

import io
import secrets

import pytest
from docx import Document

from app.models.entities import DetectedEntity
from app.services.docx_deanonymize import (
    DeanonymizeResult,
    PlaceholderMatcher,
    deanonymize_docx,
    scan_placeholders,
)
from app.services.entity_registry import EntityRegistry


# ── Helpers ───────────────────────────────────────────────────────


def _make_registry(
    locale: str = "ru",
    entities: list[tuple[str, str]] | None = None,
) -> EntityRegistry:
    """Create a registry pre-populated with entities.

    ``entities`` is a list of (text, entity_type) pairs.
    """
    key = secrets.token_bytes(32)
    reg = EntityRegistry(key, "test-session", locale=locale)
    for text, etype in (entities or []):
        e = DetectedEntity(
            text=text,
            entity_type=etype,
            start=0,
            end=len(text),
            score=0.95,
        )
        reg.get_or_create_placeholder(e)
    return reg


def _make_docx(paragraphs: list[str]) -> bytes:
    """Create a minimal .docx in memory with the given paragraph texts."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PlaceholderMatcher unit tests ─────────────────────────────────


class TestPlaceholderMatcherExact:
    """Exact match — placeholder format identical to registry."""

    def test_exact_ru(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        norm, val = matcher.match("[ЛИЦО_1]")
        assert norm == "[ЛИЦО_1]"
        # BUG-7 fix: match() returns original_forms[0], not canonical_value
        entry = reg._reverse["[ЛИЦО_1]"]
        assert val == (entry.original_forms[0] if entry.original_forms else entry.canonical_value)

    def test_exact_org(self):
        reg = _make_registry(entities=[("ООО Ромашка", "ORG")])
        matcher = PlaceholderMatcher(reg)
        norm, val = matcher.match("[ОРГАНИЗАЦИЯ_1]")
        assert norm == "[ОРГАНИЗАЦИЯ_1]"
        assert val is not None

    def test_exact_en_locale(self):
        reg = _make_registry(locale="en", entities=[("John Doe", "PER")])
        matcher = PlaceholderMatcher(reg)
        norm, val = matcher.match("[PERSON_1]")
        assert norm == "[PERSON_1]"
        assert val is not None


class TestPlaceholderMatcherFuzzy:
    """Fuzzy normalization — LLM distortions of placeholder format."""

    def test_space_instead_of_underscore(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[ЛИЦО 1]")
        assert val is not None

    def test_missing_brackets(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("ЛИЦО_1")
        assert val is not None

    def test_extra_spaces(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[ ЛИЦО_1 ]")
        assert val is not None

    def test_lowercase(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[лицо_1]")
        assert val is not None

    def test_dash_instead_of_underscore(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[ЛИЦО-1]")
        assert val is not None

    def test_leading_zero(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[ЛИЦО_01]")
        assert val is not None

    def test_cross_language_en_to_ru(self):
        """LLM writes English label but registry is Russian."""
        reg = _make_registry(locale="ru", entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[PERSON_1]")
        assert val is not None

    def test_cross_language_ru_to_en(self):
        """LLM writes Russian label but registry is English."""
        reg = _make_registry(locale="en", entities=[("John", "PER")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[ЛИЦО_1]")
        assert val is not None


class TestPlaceholderMatcherLabelCollision:
    """BUG-P2-1 regression: custom 4-letter Cyrillic labels must stay
    distinct from near-neighbour built-in labels.

    Before the length-proportional threshold, ``АФТА`` (Levenshtein 2
    from ``ДАТА``) silently resolved to a ДАТА entry with the same
    number, corrupting the round-trip with a date string.
    """

    def test_custom_afta_does_not_collapse_to_data(self):
        # Registry has DATE in slot 2; the response doc uses a *different*
        # custom label АФТА that the user introduced, also in slot 2.
        reg = _make_registry(
            entities=[
                ("10.01.2026", "ДАТА"),
                ("2026-01-10", "ДАТА"),  # now [ДАТА_2]
                ("какое-то значение", "АФТА"),  # now [АФТА_1]
                ("второе значение АФТА", "АФТА"),  # now [АФТА_2]
            ],
        )
        matcher = PlaceholderMatcher(reg)
        norm, val = matcher.match("[АФТА_2]")
        assert norm == "[АФТА_2]"
        assert val == "второе значение АФТА"

    def test_unknown_4letter_label_stays_unresolved(self):
        # Registry only has ДАТА, doc contains a 4-letter placeholder
        # that is Levenshtein-2 away. Must NOT match — must be unresolved.
        reg = _make_registry(entities=[("01.01.2026", "ДАТА")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[АФТА_1]")
        assert val is None

    def test_ocr_style_single_char_typo_still_matches(self):
        # Guard that the new threshold does not over-block: a single
        # character typo in a 4+ letter label should still resolve.
        reg = _make_registry(entities=[("Москва", "МЕСТО")])
        matcher = PlaceholderMatcher(reg)
        # МЕСТО vs МЕСТ0: 1-char substitution in a 5-char label → accept.
        _, val = matcher.match("[МЕСТ0_1]")
        assert val is not None

    def test_3letter_labels_require_exact_match(self):
        # Below 4 chars the threshold drops to 0: any single edit is a
        # 25%+ deformation and too risky.
        reg = _make_registry(entities=[("Иванов", "ЛИЦ")])
        matcher = PlaceholderMatcher(reg)
        _, val = matcher.match("[ЛИС_1]")
        assert val is None


class TestPlaceholderMatcherUnresolved:
    """Unresolved — placeholder not in registry."""

    def test_unknown_number(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        norm, val = matcher.match("[ЛИЦО_5]")
        assert val is None
        assert norm == "[ЛИЦО_5]"

    def test_unknown_type(self):
        reg = _make_registry(entities=[("Иванов", "PER")])
        matcher = PlaceholderMatcher(reg)
        # "МЕСТО_1" — there's no LOC entity in registry
        norm, val = matcher.match("[МЕСТО_1]")
        assert val is None


# ── scan_placeholders ─────────────────────────────────────────────


class TestScanPlaceholders:
    def test_finds_standard(self):
        text = "Договор между [ЛИЦО_1] и [ОРГАНИЗАЦИЯ_2] от [ДАТА_1]."
        matches = scan_placeholders(text)
        assert len(matches) == 3

    def test_finds_distorted(self):
        text = "Ответ: ЛИЦО_1 сообщил [ЛИЦО 2] что [Person_3] придёт."
        matches = scan_placeholders(text)
        assert len(matches) >= 3

    def test_empty_text(self):
        assert scan_placeholders("") == []


# ── deanonymize_docx integration ─────────────────────────────────


class TestDeanonymizeDocx:
    def test_basic_round_trip(self):
        """Simple document with two placeholders → both replaced."""
        reg = _make_registry(
            entities=[("Иванов Иван", "PER"), ("ООО Ромашка", "ORG")],
        )
        docx_bytes = _make_docx([
            "Настоящий договор заключён между [ЛИЦО_1] и [ОРГАНИЗАЦИЯ_1].",
            "Стороны договорились о нижеследующем.",
        ])

        result = deanonymize_docx(docx_bytes, reg)

        assert isinstance(result, DeanonymizeResult)
        assert len(result.replacements) >= 2
        assert len(result.unresolved) == 0

        # Verify the output DOCX contains the real values
        doc = Document(io.BytesIO(result.docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        # BUG-7 fix: deanonymize returns original_forms[0], not canonical
        per_entry = reg._reverse["[ЛИЦО_1]"]
        org_entry = reg._reverse["[ОРГАНИЗАЦИЯ_1]"]
        per_val = per_entry.original_forms[0] if per_entry.original_forms else per_entry.canonical_value
        org_val = org_entry.original_forms[0] if org_entry.original_forms else org_entry.canonical_value
        assert per_val in full_text
        assert org_val in full_text
        assert "[ЛИЦО_1]" not in full_text
        assert "[ОРГАНИЗАЦИЯ_1]" not in full_text

    def test_unresolved_reported(self):
        """Placeholder not in registry → appears in unresolved list."""
        reg = _make_registry(entities=[("Иванов", "PER")])
        docx_bytes = _make_docx([
            "[ЛИЦО_1] и [ЛИЦО_5] подписали договор.",
        ])

        result = deanonymize_docx(docx_bytes, reg)
        assert len(result.unresolved) >= 1
        norms = [u.normalized for u in result.unresolved]
        assert "[ЛИЦО_5]" in norms

    def test_manual_resolutions(self):
        """Manual resolution fills in unresolved placeholders."""
        reg = _make_registry(entities=[("Иванов", "PER")])
        docx_bytes = _make_docx(["[ЛИЦО_1] и [ЛИЦО_5] подписали."])

        result = deanonymize_docx(
            docx_bytes,
            reg,
            manual_resolutions={"[ЛИЦО_5]": "Петров Пётр"},
        )

        doc = Document(io.BytesIO(result.docx_bytes))
        text = doc.paragraphs[0].text
        assert "Петров Пётр" in text
        assert "[ЛИЦО_5]" not in text

    def test_manual_resolution_for_custom_placeholder_with_fuzzy_key(self):
        reg = _make_registry(entities=[("\u0418\u0432\u0430\u043d\u043e\u0432", "PER")])
        docx_bytes = _make_docx(["\u0414\u043e\u0433\u043e\u0432\u043e\u0440 \u0434\u0435\u0439\u0441\u0442\u0432\u0443\u0435\u0442 [\u0421\u0420\u041e\u041a_1] \u043b\u0435\u0442."])

        result = deanonymize_docx(
            docx_bytes,
            reg,
            manual_resolutions={"\u0421\u0420\u041e\u041a 1": "25"},
        )

        doc = Document(io.BytesIO(result.docx_bytes))
        text = doc.paragraphs[0].text
        assert "25" in text
        assert "[\u0421\u0420\u041e\u041a_1]" not in text
        assert len(result.unresolved) == 0

    def test_fuzzy_matching_in_docx(self):
        """LLM distortions are resolved in actual DOCX replacement."""
        reg = _make_registry(entities=[("Иванов", "PER")])
        docx_bytes = _make_docx(["Ответил [ЛИЦО 1] на запрос."])

        result = deanonymize_docx(docx_bytes, reg)

        doc = Document(io.BytesIO(result.docx_bytes))
        text = doc.paragraphs[0].text
        per_entry = reg._reverse["[ЛИЦО_1]"]
        per_val = per_entry.original_forms[0] if per_entry.original_forms else per_entry.canonical_value
        assert per_val in text
        assert len(result.unresolved) == 0

    def test_empty_document(self):
        """Empty document returns empty result without errors."""
        reg = _make_registry(entities=[("Иванов", "PER")])
        docx_bytes = _make_docx([""])

        result = deanonymize_docx(docx_bytes, reg)
        assert len(result.replacements) == 0
        assert len(result.unresolved) == 0

    def test_multiple_same_placeholder(self):
        """Same placeholder appearing multiple times → all replaced."""
        reg = _make_registry(entities=[("Иванов", "PER")])
        docx_bytes = _make_docx([
            "[ЛИЦО_1] подписал. Затем [ЛИЦО_1] подтвердил.",
        ])

        result = deanonymize_docx(docx_bytes, reg)
        doc = Document(io.BytesIO(result.docx_bytes))
        text = doc.paragraphs[0].text
        assert text.count("[ЛИЦО_1]") == 0
