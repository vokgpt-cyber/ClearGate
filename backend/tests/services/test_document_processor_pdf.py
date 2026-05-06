"""PDF/TXT document ingestion regression tests."""

from __future__ import annotations

import io

from docx import Document

from app.services.doc_processor import DocumentProcessor


def _sample_pdf_bytes() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 90), "DISTRIBUTION")
    page.insert_text((72, 120), "No. D-15-2024")
    page.insert_text((72, 150), "BETWEEN: AO Nord-Him / Tianjin Forward Polymers Co.")
    page.insert_text((72, 190), "Price is 4 500 000 CNY, including VAT 20%.")
    data = doc.tobytes()
    doc.close()
    return data


def test_pdf_parse_extracts_text_and_builds_docx_preview() -> None:
    result = DocumentProcessor().parse(_sample_pdf_bytes(), "pdf")

    assert result.page_count == 1
    assert "Tianjin Forward Polymers Co." in result.text
    assert "4 500 000 CNY" in result.text
    assert result.render_docx_bytes
    assert result.render_filename == "document_from_pdf.docx"

    doc = Document(io.BytesIO(result.render_docx_bytes))
    preview_text = "\n".join(p.text for p in doc.paragraphs)
    assert "DISTRIBUTION" in preview_text
    assert "Tianjin Forward Polymers Co." in preview_text


def test_txt_parse_builds_docx_preview() -> None:
    result = DocumentProcessor().parse(
        "Contract\nAO Nord-Him\nINN 7801456328".encode("utf-8"),
        "txt",
    )

    assert result.page_count is None
    assert "AO Nord-Him" in result.text
    assert result.render_docx_bytes

    doc = Document(io.BytesIO(result.render_docx_bytes))
    preview_text = "\n".join(p.text for p in doc.paragraphs)
    assert "INN 7801456328" in preview_text


def test_pdf_text_normalizer_repairs_legal_entity_line_breaks() -> None:
    raw = (
        "Заключен между ООО «Промышленная\n"
        "недвижимость СПб» и АО «Норд-Хим».\n"
        "Адрес: СПб, ул. Литераторов, д. 20 Банк: ПАО ВТБ\n"
    )

    normalized = DocumentProcessor()._normalize_pdf_text(raw)

    assert "ООО «Промышленная недвижимость СПб»" in normalized
    assert "д. 20\nБанк: ПАО ВТБ" in normalized
