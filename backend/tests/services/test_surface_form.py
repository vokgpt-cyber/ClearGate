"""Regression tests for BUG-P2-2: case-form (surface-form) preservation.

These tests verify that inflected Russian words survive the full
anonymize → deanonymize round-trip without collapsing to the lemma.

Layer 1: deanonymize_text uses original_forms[0] not canonical_value.
Layer 2: deanonymize_docx uses per-occurrence surface_forms.
"""

from __future__ import annotations

import io
import secrets

import pytest
from docx import Document

from app.models.entities import DetectedEntity
from app.services.entity_registry import EntityRegistry, MappingEntry
from app.services.docx_deanonymize import (
    PlaceholderMatcher,
    deanonymize_docx,
)


def _entity(text, entity_type="LOC", start=0, end=None, score=0.9):
    return DetectedEntity(
        text=text,
        entity_type=entity_type,
        start=start,
        end=end or start + len(text),
        score=score,
    )


@pytest.fixture()
def master_key():
    return secrets.token_bytes(32)


@pytest.fixture()
def registry(master_key):
    return EntityRegistry(master_key, "surface-test", locale="ru")


# ═══════════════════════════════════════════════════════════════════
# Layer 1 — deanonymize_text
# ═══════════════════════════════════════════════════════════════════


class TestLayer1DeanonymizeText:
    """deanonymize_text should return original_forms[0], not the lemma."""

    def test_single_inflected_form_preserved(self, registry):
        """'Москве' (prepositional) must NOT become 'Москва' (nominative)."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        result = registry.deanonymize_text("[МЕСТО_1] — прекрасный город.")
        assert "Москве" in result
        assert "москва" not in result.lower().replace("москве", "")

    def test_person_name_inflection_preserved(self, registry):
        """'Иванову' (dative) must NOT become lemma."""
        registry.get_or_create_placeholder(_entity("Иванову", "PER"))
        result = registry.deanonymize_text("Передать документы [ЛИЦО_1].")
        assert "Иванову" in result

    def test_empty_original_forms_falls_back_to_canonical(self, registry):
        """Safety: if original_forms is empty, canonical_value is used."""
        registry.get_or_create_placeholder(_entity("Иванов", "PER"))
        # Simulate an edge case: clear original_forms manually
        entry = list(registry._reverse.values())[0]
        entry.original_forms.clear()
        result = registry.deanonymize_text("[ЛИЦО_1]")
        # Should use canonical_value (the lemma) as fallback
        assert "[ЛИЦО_1]" not in result


# ═══════════════════════════════════════════════════════════════════
# Layer 2 — surface_forms + occurrence-aware deanonymize_docx
# ═══════════════════════════════════════════════════════════════════


class TestRecordSurfaceForms:
    """EntityRegistry.record_surface_forms builds correct per-occurrence list."""

    def test_basic_recording(self, registry):
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.get_or_create_placeholder(_entity("Москва", "LOC", start=20))
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            ("Москва", "[МЕСТО_1]"),
        ])
        entry = registry._reverse["[МЕСТО_1]"]
        assert entry.surface_forms == ["Москве", "Москва"]

    def test_three_forms(self, registry):
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.get_or_create_placeholder(_entity("Москва", "LOC", start=20))
        registry.get_or_create_placeholder(_entity("Москвы", "LOC", start=40))
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            ("Москва", "[МЕСТО_1]"),
            ("Москвы", "[МЕСТО_1]"),
        ])
        entry = registry._reverse["[МЕСТО_1]"]
        assert entry.surface_forms == ["Москве", "Москва", "Москвы"]

    def test_recording_clears_previous(self, registry):
        """Calling record_surface_forms twice replaces, doesn't append."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.record_surface_forms([("Москве", "[МЕСТО_1]")])
        registry.record_surface_forms([("Москва", "[МЕСТО_1]")])
        entry = registry._reverse["[МЕСТО_1]"]
        assert entry.surface_forms == ["Москва"]

    def test_unknown_placeholder_ignored(self, registry):
        """Placeholders not in the registry are silently skipped."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            ("Foo", "[НЕСУЩЕСТВУЮЩИЙ_99]"),
        ])
        entry = registry._reverse["[МЕСТО_1]"]
        assert entry.surface_forms == ["Москве"]

    def test_rejected_entities_excluded(self, registry):
        """Only accepted (non-rejected) entities should be passed in."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.get_or_create_placeholder(_entity("Москва", "LOC", start=20))
        registry.get_or_create_placeholder(_entity("Москвы", "LOC", start=40))
        # Simulating: user rejected the 2nd occurrence ("Москва")
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            # "Москва" omitted — rejected
            ("Москвы", "[МЕСТО_1]"),
        ])
        entry = registry._reverse["[МЕСТО_1]"]
        assert entry.surface_forms == ["Москве", "Москвы"]


class TestPlaceholderMatcherOccurrenceAware:
    """PlaceholderMatcher._get_real_value respects occurrence_index."""

    def test_first_occurrence_gets_first_form(self, registry):
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.get_or_create_placeholder(_entity("Москва", "LOC", start=20))
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            ("Москва", "[МЕСТО_1]"),
        ])
        matcher = PlaceholderMatcher(registry)
        _, val0 = matcher.match("[МЕСТО_1]", occurrence_index=0)
        _, val1 = matcher.match("[МЕСТО_1]", occurrence_index=1)
        assert val0 == "Москве"
        assert val1 == "Москва"

    def test_excess_occurrence_falls_back(self, registry):
        """Occurrence beyond surface_forms length → original_forms[0]."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.record_surface_forms([("Москве", "[МЕСТО_1]")])
        matcher = PlaceholderMatcher(registry)
        _, val = matcher.match("[МЕСТО_1]", occurrence_index=5)
        assert val == "Москве"  # falls back to original_forms[0]

    def test_empty_surface_forms_falls_back(self, registry):
        """Old sessions without surface_forms → original_forms[0]."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        # Don't call record_surface_forms — simulates an old session
        matcher = PlaceholderMatcher(registry)
        _, val = matcher.match("[МЕСТО_1]", occurrence_index=0)
        assert val == "Москве"  # falls back to original_forms[0]


class TestPersistenceBackcompat:
    """Encrypted blob round-trip with and without surface_forms."""

    def test_new_field_survives_export_import(self, registry, master_key):
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            ("Москва", "[МЕСТО_1]"),
        ])
        blob = registry.export_encrypted()

        restored = EntityRegistry(master_key, "surface-test", locale="ru")
        restored.import_encrypted(blob)
        entry = restored._reverse["[МЕСТО_1]"]
        assert entry.surface_forms == ["Москве", "Москва"]

    def test_old_blob_without_surface_forms(self, registry, master_key):
        """Blobs from Phase 1 (no surface_forms) load cleanly."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        blob = registry.export_encrypted()

        # Simulate old blob: decrypt, remove surface_forms, re-encrypt
        data = registry._crypto.decrypt_mapping(blob, "surface-test")
        for entry_data in data.get("entries", []):
            entry_data.pop("surface_forms", None)
        patched_blob = registry._crypto.encrypt_mapping(data, "surface-test")

        restored = EntityRegistry(master_key, "surface-test", locale="ru")
        restored.import_encrypted(patched_blob)
        entry = restored._reverse["[МЕСТО_1]"]
        assert entry.surface_forms == []  # Pydantic default
        # Fallback to original_forms[0] should still work
        matcher = PlaceholderMatcher(restored)
        _, val = matcher.match("[МЕСТО_1]", occurrence_index=0)
        assert val == "Москве"


def _make_docx_with_placeholders(*paragraphs: str) -> bytes:
    """Create a minimal DOCX with the given paragraph texts."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDeanonymizeDocxOccurrenceAware:
    """Full deanonymize_docx with different surface forms per occurrence."""

    def test_two_forms_restored_correctly(self, registry):
        """'в [МЕСТО_1]' and '[МЕСТО_1] является' get different forms."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.get_or_create_placeholder(_entity("Москва", "LOC", start=20))
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            ("Москва", "[МЕСТО_1]"),
        ])

        docx_bytes = _make_docx_with_placeholders(
            "Договор заключен в городе [МЕСТО_1].",
            "[МЕСТО_1] является столицей.",
        )

        result = deanonymize_docx(docx_bytes, registry)
        doc = Document(io.BytesIO(result.docx_bytes))
        texts = [p.text for p in doc.paragraphs]
        assert "Москве" in texts[0], f"Expected 'Москве' in first para, got: {texts[0]}"
        assert "Москва" in texts[1], f"Expected 'Москва' in second para, got: {texts[1]}"

    def test_uniform_forms_still_work(self, registry):
        """When all occurrences have the same form, bulk replace works."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        registry.record_surface_forms([
            ("Москве", "[МЕСТО_1]"),
            ("Москве", "[МЕСТО_1]"),
        ])

        docx_bytes = _make_docx_with_placeholders(
            "в [МЕСТО_1] хорошо",
            "в [МЕСТО_1] тоже",
        )

        result = deanonymize_docx(docx_bytes, registry)
        doc = Document(io.BytesIO(result.docx_bytes))
        for p in doc.paragraphs:
            assert "Москве" in p.text
            assert "[МЕСТО_1]" not in p.text

    def test_no_surface_forms_fallback(self, registry):
        """Old sessions without surface_forms still deanonymize correctly."""
        registry.get_or_create_placeholder(_entity("Москве", "LOC"))
        # No record_surface_forms call

        docx_bytes = _make_docx_with_placeholders("в городе [МЕСТО_1]")
        result = deanonymize_docx(docx_bytes, registry)
        doc = Document(io.BytesIO(result.docx_bytes))
        assert "Москве" in doc.paragraphs[0].text
